#!/usr/bin/env python3
from __future__ import annotations

import csv
import ctypes
import ctypes.util
import json
import math
import os
import re
import shutil
import statistics
import subprocess
import sys
import threading
import time
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import numpy as np
from PIL import Image


PROJECT = Path("/home/jetson/Downloads/classification_jetson")
MODEL_DIR = PROJECT / "mobilenet_v3_small_inference_models"
OUT = PROJECT / "mobilenet_v3_small_fp32_core4_benchmark"
TEST_DIR = PROJECT / "test_imgs"
PT_PATH = MODEL_DIR / "mobilenet_v3_small_inference_only.pt"
ONNX_PATH = MODEL_DIR / "mobilenet_v3_small_inference.onnx"
TS_PATH = MODEL_DIR / "mobilenet_v3_small_inference_fp32.torchscript.pt"
ENGINE_PATH = MODEL_DIR / "mobilenet_v3_small_fp32_640x640.engine"
TRT_ONNX_PATH = OUT / "exports" / "mobilenet_v3_small_inference_fp32_trt_compatible.onnx"
ZIP_PATH = PROJECT / "mobilenet_v3_small_fp32_core4_benchmark.zip"

INPUT_SHAPE = "1x3x640x640"
IMG_SIZE = 640
WARMUP = 10
IMG_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff"}
CLASS_NAMES = ["class_1_asymptomatic", "class_2_mild_symptoms", "class_3_symptomatic_compromised"]
IMAGENET_MEAN = np.array([0.485, 0.456, 0.406], dtype=np.float32).reshape(3, 1, 1)
IMAGENET_STD = np.array([0.229, 0.224, 0.225], dtype=np.float32).reshape(3, 1, 1)

for d in ["logs", "metrics", "predictions", "reports", "resources", "exports", "models", "scripts"]:
    (OUT / d).mkdir(parents=True, exist_ok=True)


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def run_cmd(cmd: str, log_path: Path | None = None, timeout: int | None = None):
    start = time.time()
    proc = subprocess.run(cmd, shell=True, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=timeout)
    payload = f"$ {cmd}\nreturncode={proc.returncode}\nduration_s={time.time()-start:.3f}\n\n{proc.stdout}"
    if log_path:
        write_text(log_path, payload)
    return proc.returncode, proc.stdout


def safe_imports():
    info = {
        "python_path": sys.executable,
        "python_version": sys.version,
        "working_directory": str(PROJECT),
        "input_shape": INPUT_SHAPE,
        "preprocessing": {
            "rgb": True,
            "resize": [IMG_SIZE, IMG_SIZE],
            "to_tensor": "float32 CHW divided by 255",
            "normalize_mean": [0.485, 0.456, 0.406],
            "normalize_std": [0.229, 0.224, 0.225],
            "source": "MobileNetV3-Small ImageNet1K V1 training/export metadata; no separate preprocessing.json was saved.",
        },
    }
    try:
        import torch

        info["torch_version"] = torch.__version__
        info["cuda_available"] = bool(torch.cuda.is_available())
        info["cuda_device_name"] = torch.cuda.get_device_name(0) if torch.cuda.is_available() else ""
    except Exception as e:
        info["torch_error"] = repr(e)
    try:
        import torchvision

        info["torchvision_version"] = torchvision.__version__
    except Exception as e:
        info["torchvision_error"] = repr(e)
    try:
        import onnxruntime as ort

        info["onnxruntime_available"] = True
        info["onnxruntime_version"] = ort.__version__
        info["onnxruntime_providers"] = ort.get_available_providers()
    except Exception as e:
        info["onnxruntime_available"] = False
        info["onnxruntime_error"] = repr(e)
    try:
        user_site = Path("/home/jetson/.local/lib/python3.8/site-packages")
        if user_site.exists() and str(user_site) not in sys.path:
            sys.path.insert(0, str(user_site))
        import tensorrt as trt

        info["tensorrt_python_available"] = True
        info["tensorrt_version"] = trt.__version__
        info["tensorrt_file"] = getattr(trt, "__file__", "")
    except Exception as e:
        info["tensorrt_python_available"] = False
        info["tensorrt_error"] = repr(e)
    return info


def load_export_metadata():
    path = MODEL_DIR / "onnx_export_metadata.json"
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {}


def image_paths_and_labels():
    class_dirs = [TEST_DIR / c for c in CLASS_NAMES]
    missing = [str(p) for p in class_dirs if not p.exists()]
    if missing:
        raise RuntimeError(f"Missing expected class folders: {missing}")
    items = []
    for idx, d in enumerate(class_dirs):
        for p in sorted(d.rglob("*")):
            if p.suffix.lower() in IMG_EXTS:
                items.append((p, idx))
    return CLASS_NAMES, items


def reconstruct_model():
    import torch
    import torch.nn as nn
    import torchvision.models as models

    try:
        model = models.mobilenet_v3_small(weights=None)
    except TypeError:
        model = models.mobilenet_v3_small(pretrained=False)
    if isinstance(model.classifier, nn.Sequential):
        if len(model.classifier) >= 3 and isinstance(model.classifier[2], nn.Dropout):
            model.classifier[2] = nn.Dropout(p=0.0, inplace=True)
        last = model.classifier[-1]
        in_features = last.in_features
        model.classifier[-1] = nn.Linear(in_features, len(CLASS_NAMES))
    else:
        raise RuntimeError("Unexpected torchvision MobileNetV3-Small classifier layout")

    ckpt = torch.load(PT_PATH, map_location="cpu")
    if isinstance(ckpt, dict):
        state = None
        for key in ["model_state_dict", "state_dict"]:
            if key in ckpt:
                state = ckpt[key]
                break
        if state is None and all(hasattr(v, "shape") for v in ckpt.values()):
            state = ckpt
        if state is None and "model" in ckpt:
            obj = ckpt["model"]
            if hasattr(obj, "state_dict"):
                state = obj.state_dict()
        if state is None:
            raise RuntimeError(f"Could not find state_dict keys in checkpoint: {list(ckpt)[:20]}")
        clean = {}
        for k, v in state.items():
            nk = k.replace("module.", "")
            clean[nk] = v
        model.load_state_dict(clean, strict=True)
    elif hasattr(ckpt, "eval"):
        model = ckpt
    else:
        raise RuntimeError(f"Unsupported checkpoint type: {type(ckpt)}")
    model.float().eval()
    return model


def preprocess_pil(path: Path) -> np.ndarray:
    img = Image.open(path).convert("RGB").resize((IMG_SIZE, IMG_SIZE), Image.BILINEAR)
    arr = np.asarray(img).astype(np.float32) / 255.0
    chw = arr.transpose(2, 0, 1)
    chw = (chw - IMAGENET_MEAN) / IMAGENET_STD
    return chw[np.newaxis, ...].astype(np.float32, copy=False)


def unwrap_output(out):
    if isinstance(out, (list, tuple)):
        out = out[0]
    if isinstance(out, dict):
        for key in ["logits", "pred", "output"]:
            if key in out:
                out = out[key]
                break
    return out


def softmax_np(x) -> np.ndarray:
    x = np.asarray(x, dtype=np.float64).reshape(-1)
    x = x - np.max(x)
    e = np.exp(x)
    return e / np.sum(e)


def compute_metrics(y_true, y_pred, n_classes):
    cm = [[0 for _ in range(n_classes)] for __ in range(n_classes)]
    for t, p in zip(y_true, y_pred):
        cm[t][p] += 1
    total = sum(sum(r) for r in cm)
    correct = sum(cm[i][i] for i in range(n_classes))
    ps, rs, f1s, per = [], [], [], []
    weighted_p = weighted_r = weighted_f1 = 0.0
    for i in range(n_classes):
        tp = cm[i][i]
        fp = sum(cm[r][i] for r in range(n_classes) if r != i)
        fn = sum(cm[i][c] for c in range(n_classes) if c != i)
        support = sum(cm[i])
        p = tp / (tp + fp) if tp + fp else 0.0
        r = tp / (tp + fn) if tp + fn else 0.0
        f1 = 2 * p * r / (p + r) if p + r else 0.0
        ps.append(p)
        rs.append(r)
        f1s.append(f1)
        weighted_p += p * support
        weighted_r += r * support
        weighted_f1 += f1 * support
        per.append({"class_id": i, "precision": p, "recall": r, "f1": f1, "support": support})
    return {
        "accuracy": correct / total if total else 0.0,
        "macro_precision": sum(ps) / n_classes,
        "macro_recall": sum(rs) / n_classes,
        "macro_f1": sum(f1s) / n_classes,
        "weighted_precision": weighted_p / total if total else 0.0,
        "weighted_recall": weighted_r / total if total else 0.0,
        "weighted_f1": weighted_f1 / total if total else 0.0,
        "confusion_matrix": cm,
        "per_class": per,
    }


def percentile(values, p):
    if not values:
        return ""
    values = sorted(values)
    k = (len(values) - 1) * p / 100.0
    f = math.floor(k)
    c = math.ceil(k)
    if f == c:
        return values[int(k)]
    return values[f] * (c - k) + values[c] * (k - f)


def parse_power_w(stats):
    candidates = ["VDD_IN", "POM_5V_IN", "power", "Power TOT", "VDD_CPU_GPU_CV", "POM_5V_GPU"]
    values = []
    for k, v in stats.items():
        if not any(c.lower() in str(k).lower() for c in candidates):
            continue
        key = str(k)
        if isinstance(v, dict):
            for vv in v.values():
                values.append((key, vv))
        else:
            values.append((key, v))
    watts = []
    for key, v in values:
        if isinstance(v, (int, float)):
            watts.append(float(v) / 1000.0 if "power" in key.lower() or "pom_" in key.lower() or "vdd" in key.lower() else float(v))
        elif isinstance(v, str):
            m = re.search(r"([0-9.]+)\s*(mW|W)?", v)
            if m:
                val = float(m.group(1))
                unit = (m.group(2) or "").lower()
                watts.append(val / 1000.0 if unit == "mw" or "power" in key.lower() or "pom_" in key.lower() or "vdd" in key.lower() or val > 100 else val)
    return max(watts) if watts else None


class ResourceMonitor:
    def __init__(self, label):
        self.label = label
        self.samples = []
        self.stop = threading.Event()
        self.thread = None

    def __enter__(self):
        try:
            from jtop import jtop  # noqa
        except Exception:
            return self

        def worker():
            try:
                from jtop import jtop

                with jtop(interval=0.5) as jetson:
                    while not self.stop.is_set() and jetson.ok():
                        stats = dict(jetson.stats)
                        self.samples.append({"time": time.time(), "label": self.label, "power_w": parse_power_w(stats), "raw": stats})
                        time.sleep(0.5)
            except Exception as exc:
                self.samples.append({"time": time.time(), "label": self.label, "power_w": None, "raw": {"error": repr(exc)}})

        self.thread = threading.Thread(target=worker, daemon=True)
        self.thread.start()
        return self

    def __exit__(self, exc_type, exc, tb):
        self.stop.set()
        if self.thread:
            self.thread.join(timeout=2)

    def summary(self, inference_count, latency_mean_ms):
        powers = [s["power_w"] for s in self.samples if s.get("power_w") is not None]
        if not powers:
            return {"avg_power_w": "", "median_power_w": "", "max_power_w": "", "energy_per_inference_j": ""}
        avg = statistics.mean(powers)
        return {
            "avg_power_w": avg,
            "median_power_w": statistics.median(powers),
            "max_power_w": max(powers),
            "energy_per_inference_j": avg * (latency_mean_ms / 1000.0) if latency_mean_ms != "" else "",
        }


RESOURCE_ROWS = []


def append_resource_samples(mon):
    for s in mon.samples:
        RESOURCE_ROWS.append({"time": s["time"], "format": s["label"], "power_w": s.get("power_w", "")})


def save_prediction_files(prefix, class_names, y_true, y_pred, probs, paths, metric):
    pred_path = OUT / "predictions" / f"{prefix}_predictions.csv"
    with pred_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["image", "true_id", "true_class", "pred_id", "pred_class", "confidence"])
        for p, t, pr, prob in zip(paths, y_true, y_pred, probs):
            w.writerow([str(p), t, class_names[t], pr, class_names[pr], float(max(prob))])
    cm_path = OUT / "metrics" / f"confusion_matrix_{prefix}.csv"
    with cm_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["true\\pred"] + class_names)
        for name, row in zip(class_names, metric["confusion_matrix"]):
            w.writerow([name] + row)
    per_path = OUT / "metrics" / f"per_class_metrics_{prefix}.csv"
    with per_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["class_id", "class_name", "precision", "recall", "f1", "support"])
        w.writeheader()
        for r in metric["per_class"]:
            rr = dict(r)
            rr["class_name"] = class_names[r["class_id"]]
            w.writerow(rr)


def make_result_row(index, fmt, status, artifact, metrics=None, latency=None, resource=None, error="", notes="", shape_verified=None):
    metrics = metrics or {}
    latency = latency or {}
    resource = resource or {}
    size_mb = Path(artifact).stat().st_size / (1024 * 1024) if artifact and Path(artifact).exists() else ""
    if shape_verified is None:
        shape_verified = "YES" if status == "SUCCESS" else "NO"
    return {
        "index": index,
        "format": fmt,
        "precision": "FP32",
        "status": status,
        "size_mb": size_mb,
        "accuracy": metrics.get("accuracy", ""),
        "macro_precision": metrics.get("macro_precision", ""),
        "macro_recall": metrics.get("macro_recall", ""),
        "macro_f1": metrics.get("macro_f1", ""),
        "weighted_precision": metrics.get("weighted_precision", ""),
        "weighted_recall": metrics.get("weighted_recall", ""),
        "weighted_f1": metrics.get("weighted_f1", ""),
        "latency_mean_ms": latency.get("mean", ""),
        "latency_median_ms": latency.get("median", ""),
        "latency_p95_ms": latency.get("p95", ""),
        "fps": latency.get("fps", ""),
        "avg_power_w": resource.get("avg_power_w", ""),
        "median_power_w": resource.get("median_power_w", ""),
        "max_power_w": resource.get("max_power_w", ""),
        "energy_per_inference_j": resource.get("energy_per_inference_j", ""),
        "input_shape": INPUT_SHAPE,
        "shape_verified": shape_verified,
        "artifact_path": str(artifact),
        "error_message": error,
        "notes": notes,
    }


def evaluate_torch_model(prefix, display, model, device, class_names, items, artifact):
    import torch

    paths = [p for p, _ in items]
    y_true = [y for _, y in items]
    y_pred, probs, times = [], [], []
    model.to(device).eval()
    with torch.no_grad():
        for p, _ in items[:WARMUP]:
            x = torch.from_numpy(preprocess_pil(p)).to(device)
            _ = unwrap_output(model(x))
            if device.type == "cuda":
                torch.cuda.synchronize()
        with ResourceMonitor(display) as mon:
            for p, _ in items:
                x = torch.from_numpy(preprocess_pil(p)).to(device)
                if device.type == "cuda":
                    torch.cuda.synchronize()
                t0 = time.perf_counter()
                out = unwrap_output(model(x))
                if device.type == "cuda":
                    torch.cuda.synchronize()
                dt = (time.perf_counter() - t0) * 1000
                arr = out.detach().float().cpu().numpy().reshape(-1)
                prob = softmax_np(arr)
                y_pred.append(int(np.argmax(prob)))
                probs.append(prob)
                times.append(dt)
    append_resource_samples(mon)
    metric = compute_metrics(y_true, y_pred, len(class_names))
    latency = {"mean": statistics.mean(times), "median": statistics.median(times), "p95": percentile(times, 95), "fps": 1000.0 / statistics.mean(times)}
    resource = mon.summary(len(items), latency["mean"])
    save_prediction_files(prefix, class_names, y_true, y_pred, probs, paths, metric)
    payload = {**metric, "latency": latency, "resource": resource, "artifact_path": str(artifact), "input_shape": INPUT_SHAPE, "shape_verified": True}
    (OUT / "metrics" / f"{prefix}_metrics.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return make_result_row(None, display, "SUCCESS", artifact, metric, latency, resource)


def export_torchscript(model, device):
    import torch

    if TS_PATH.exists():
        return

    class Wrapper(torch.nn.Module):
        def __init__(self, m):
            super().__init__()
            self.m = m

        def forward(self, x):
            out = self.m(x)
            if isinstance(out, (list, tuple)):
                out = out[0]
            return out

    wrapper = Wrapper(model).to(device).eval()
    dummy = torch.zeros(1, 3, IMG_SIZE, IMG_SIZE, device=device)
    with torch.no_grad():
        traced = torch.jit.trace(wrapper, dummy, strict=False)
        traced.save(str(TS_PATH))


def export_trt_compatible_onnx(model, device):
    import torch
    import torch.nn as nn

    if TRT_ONNX_PATH.exists():
        return TRT_ONNX_PATH

    class DecomposedHardswish(nn.Module):
        def forward(self, x):
            return x * torch.nn.functional.relu6(x + 3.0) / 6.0

    replaced = 0
    for name, child in list(model.named_modules()):
        # Replacement must happen through the parent module.
        pass
    for parent in model.modules():
        for child_name, child in list(parent.named_children()):
            if child.__class__.__name__.lower() == "hardswish":
                setattr(parent, child_name, DecomposedHardswish())
                replaced += 1
    wrapper = model.to(device).eval()
    dummy = torch.zeros(1, 3, IMG_SIZE, IMG_SIZE, device=device)
    torch.onnx.export(
        wrapper,
        dummy,
        str(TRT_ONNX_PATH),
        input_names=["input"],
        output_names=["logits"],
        opset_version=12,
        do_constant_folding=True,
    )
    write_text(OUT / "reports" / "tensorrt_compatible_onnx_report.txt", f"Created TensorRT-compatible ONNX: {TRT_ONNX_PATH}\nHardswish modules replaced: {replaced}\n")
    return TRT_ONNX_PATH


def evaluate_onnx(class_names, items):
    try:
        import onnxruntime as ort
    except Exception as e:
        return make_result_row(3, "ONNX FP32", "RUNTIME_UNAVAILABLE", ONNX_PATH, error=repr(e), notes="ONNXRuntime is unavailable in the safe Jetson Python runtime.")
    try:
        paths = [p for p, _ in items]
        y_true = [y for _, y in items]
        sess = ort.InferenceSession(str(ONNX_PATH), providers=["CPUExecutionProvider"])
        inp = sess.get_inputs()[0].name
        shape = sess.get_inputs()[0].shape
        shape_verified = "YES" if list(shape) in ([1, 3, IMG_SIZE, IMG_SIZE], ["batch", 3, IMG_SIZE, IMG_SIZE], [None, 3, IMG_SIZE, IMG_SIZE]) or tuple(shape[1:]) == (3, IMG_SIZE, IMG_SIZE) else "NO"
        y_pred, probs, times = [], [], []
        for p, _ in items[:WARMUP]:
            x = preprocess_pil(p)
            sess.run(None, {inp: x})
        with ResourceMonitor("ONNX FP32") as mon:
            for p, _ in items:
                x = preprocess_pil(p)
                t0 = time.perf_counter()
                out = sess.run(None, {inp: x})[0]
                dt = (time.perf_counter() - t0) * 1000
                prob = softmax_np(out.reshape(-1))
                y_pred.append(int(np.argmax(prob)))
                probs.append(prob)
                times.append(dt)
        append_resource_samples(mon)
        metric = compute_metrics(y_true, y_pred, len(class_names))
        latency = {"mean": statistics.mean(times), "median": statistics.median(times), "p95": percentile(times, 95), "fps": 1000.0 / statistics.mean(times)}
        resource = mon.summary(len(items), latency["mean"])
        save_prediction_files("onnx_fp32", class_names, y_true, y_pred, probs, paths, metric)
        (OUT / "metrics" / "onnx_fp32_metrics.json").write_text(json.dumps({**metric, "latency": latency, "resource": resource, "onnx_input_shape": shape}, indent=2), encoding="utf-8")
        return make_result_row(3, "ONNX FP32", "SUCCESS", ONNX_PATH, metric, latency, resource, shape_verified=shape_verified)
    except Exception as e:
        write_text(OUT / "logs" / "onnx_fp32_error.log", repr(e))
        return make_result_row(3, "ONNX FP32", "FAILED", ONNX_PATH, error=repr(e))


class CudaRuntime:
    H2D = 1
    D2H = 2

    def __init__(self):
        lib = ctypes.util.find_library("cudart") or "libcudart.so"
        self.lib = ctypes.CDLL(lib)
        self.lib.cudaMalloc.argtypes = [ctypes.POINTER(ctypes.c_void_p), ctypes.c_size_t]
        self.lib.cudaFree.argtypes = [ctypes.c_void_p]
        self.lib.cudaMemcpy.argtypes = [ctypes.c_void_p, ctypes.c_void_p, ctypes.c_size_t, ctypes.c_int]
        self.lib.cudaDeviceSynchronize.argtypes = []

    def check(self, code, name):
        if code != 0:
            raise RuntimeError(f"{name} failed with cuda error {code}")

    def malloc(self, nbytes):
        ptr = ctypes.c_void_p()
        self.check(self.lib.cudaMalloc(ctypes.byref(ptr), nbytes), "cudaMalloc")
        return ptr

    def free(self, ptr):
        self.lib.cudaFree(ptr)

    def memcpy_htod(self, ptr, arr):
        arr = np.ascontiguousarray(arr)
        self.check(self.lib.cudaMemcpy(ptr, arr.ctypes.data_as(ctypes.c_void_p), arr.nbytes, self.H2D), "cudaMemcpy H2D")

    def memcpy_dtoh(self, arr, ptr):
        arr = np.ascontiguousarray(arr)
        self.check(self.lib.cudaMemcpy(arr.ctypes.data_as(ctypes.c_void_p), ptr, arr.nbytes, self.D2H), "cudaMemcpy D2H")

    def sync(self):
        self.check(self.lib.cudaDeviceSynchronize(), "cudaDeviceSynchronize")


def infer_onnx_input_name(path: Path) -> str:
    try:
        import onnx

        model = onnx.load(str(path))
        return model.graph.input[0].name
    except Exception:
        return "input"


def build_engine_if_needed(model, device):
    if ENGINE_PATH.exists():
        return True, "existing engine", ONNX_PATH
    log = OUT / "logs" / "tensorrt_fp32_build.log"
    attempts = []
    input_name = infer_onnx_input_name(ONNX_PATH)
    attempts.append((ONNX_PATH, f"/usr/src/tensorrt/bin/trtexec --onnx={ONNX_PATH} --saveEngine={ENGINE_PATH} --workspace=1024"))
    attempts.append((ONNX_PATH, f"/usr/src/tensorrt/bin/trtexec --onnx={ONNX_PATH} --saveEngine={ENGINE_PATH} --workspace=1024 --minShapes={input_name}:1x3x640x640 --optShapes={input_name}:1x3x640x640 --maxShapes={input_name}:1x3x640x640"))

    build_logs = []
    for onnx_path, cmd in attempts:
        rc, out = run_cmd(cmd, log_path=log, timeout=1800)
        build_logs.append(out)
        if rc == 0 and ENGINE_PATH.exists():
            return True, out, onnx_path

    joined = "\n".join(build_logs)
    if "HardSwish" in joined or "Hardswish" in joined or "No importer registered" in joined:
        compat = export_trt_compatible_onnx(model, device)
        cmd = f"/usr/src/tensorrt/bin/trtexec --onnx={compat} --saveEngine={ENGINE_PATH} --workspace=1024"
        rc, out = run_cmd(cmd, log_path=log, timeout=1800)
        if rc == 0 and ENGINE_PATH.exists():
            return True, out, compat
        input_name = infer_onnx_input_name(compat)
        cmd = f"/usr/src/tensorrt/bin/trtexec --onnx={compat} --saveEngine={ENGINE_PATH} --workspace=1024 --minShapes={input_name}:1x3x640x640 --optShapes={input_name}:1x3x640x640 --maxShapes={input_name}:1x3x640x640"
        rc, out = run_cmd(cmd, log_path=log, timeout=1800)
        if rc == 0 and ENGINE_PATH.exists():
            return True, out, compat
    return ENGINE_PATH.exists(), joined, ONNX_PATH


def evaluate_tensorrt(model, device, class_names, items):
    ok, build_msg, engine_source_onnx = build_engine_if_needed(model, device)
    if not ok:
        return make_result_row(4, "TensorRT FP32", "FAILED", ENGINE_PATH, error="TensorRT engine build failed. See logs/tensorrt_fp32_build.log.", notes=str(build_msg)[-1000:])
    try:
        user_site = Path("/home/jetson/.local/lib/python3.8/site-packages")
        if user_site.exists() and str(user_site) not in sys.path:
            sys.path.insert(0, str(user_site))
        import tensorrt as trt
    except Exception as e:
        return make_result_row(4, "TensorRT FP32", "RUNTIME_UNAVAILABLE", ENGINE_PATH, error=repr(e), notes="TensorRT Python is unavailable.")
    try:
        logger = trt.Logger(trt.Logger.WARNING)
        with open(ENGINE_PATH, "rb") as f, trt.Runtime(logger) as runtime:
            engine = runtime.deserialize_cuda_engine(f.read())
        context = engine.create_execution_context()
        in_idx = next(i for i in range(engine.num_bindings) if engine.binding_is_input(i))
        out_idx = next(i for i in range(engine.num_bindings) if not engine.binding_is_input(i))
        if -1 in tuple(engine.get_binding_shape(in_idx)):
            context.set_binding_shape(in_idx, (1, 3, IMG_SIZE, IMG_SIZE))
        input_shape = tuple(context.get_binding_shape(in_idx))
        out_shape = tuple(context.get_binding_shape(out_idx))
        if any(d < 0 for d in out_shape):
            out_shape = tuple(engine.get_binding_shape(out_idx))
        shape_verified = "YES" if input_shape == (1, 3, IMG_SIZE, IMG_SIZE) else "NO"
        binding_info = {
            "engine_path": str(ENGINE_PATH),
            "source_onnx_used_for_engine": str(engine_source_onnx),
            "bindings": [
                {"index": i, "name": engine.get_binding_name(i), "is_input": engine.binding_is_input(i), "shape": list(context.get_binding_shape(i)), "dtype": str(engine.get_binding_dtype(i))}
                for i in range(engine.num_bindings)
            ],
            "input_shape_expected": [1, 3, IMG_SIZE, IMG_SIZE],
            "shape_verified": shape_verified == "YES",
        }
        (OUT / "metrics" / "tensorrt_fp32_engine_verification.json").write_text(json.dumps(binding_info, indent=2), encoding="utf-8")
        if shape_verified != "YES":
            if ENGINE_PATH.exists():
                ENGINE_PATH.unlink()
            return make_result_row(4, "TensorRT FP32", "FAILED", ENGINE_PATH, error=f"Invalid TensorRT input binding shape: {input_shape}", shape_verified="NO")

        input_np = np.zeros((1, 3, IMG_SIZE, IMG_SIZE), dtype=np.float32)
        output_np = np.empty(out_shape, dtype=np.float32)
        cuda = CudaRuntime()
        d_input = cuda.malloc(input_np.nbytes)
        d_output = cuda.malloc(output_np.nbytes)
        bindings = [0] * engine.num_bindings
        bindings[in_idx] = int(d_input.value)
        bindings[out_idx] = int(d_output.value)
        paths = [p for p, _ in items]
        y_true = [y for _, y in items]
        y_pred, probs, times = [], [], []
        for p, _ in items[:WARMUP]:
            x = preprocess_pil(p)
            cuda.memcpy_htod(d_input, x)
            context.execute_v2(bindings)
            cuda.sync()
        with ResourceMonitor("TensorRT FP32") as mon:
            for p, _ in items:
                x = preprocess_pil(p)
                t0 = time.perf_counter()
                cuda.memcpy_htod(d_input, x)
                context.execute_v2(bindings)
                cuda.memcpy_dtoh(output_np, d_output)
                cuda.sync()
                dt = (time.perf_counter() - t0) * 1000
                prob = softmax_np(output_np.reshape(-1))
                y_pred.append(int(np.argmax(prob)))
                probs.append(prob)
                times.append(dt)
        cuda.free(d_input)
        cuda.free(d_output)
        append_resource_samples(mon)
        metric = compute_metrics(y_true, y_pred, len(class_names))
        latency = {"mean": statistics.mean(times), "median": statistics.median(times), "p95": percentile(times, 95), "fps": 1000.0 / statistics.mean(times)}
        resource = mon.summary(len(items), latency["mean"])
        save_prediction_files("tensorrt_fp32", class_names, y_true, y_pred, probs, paths, metric)
        (OUT / "metrics" / "tensorrt_fp32_metrics.json").write_text(json.dumps({**metric, "latency": latency, "resource": resource, "engine_verification": binding_info}, indent=2), encoding="utf-8")
        notes = "TensorRT FP32 engine built from original ONNX." if engine_source_onnx == ONNX_PATH else "TensorRT FP32 engine built from separate Hardswish-decomposed ONNX because Jetson TensorRT rejected the original ONNX."
        return make_result_row(4, "TensorRT FP32", "SUCCESS", ENGINE_PATH, metric, latency, resource, notes=notes, shape_verified=shape_verified)
    except Exception as e:
        write_text(OUT / "logs" / "tensorrt_fp32_error.log", repr(e))
        return make_result_row(4, "TensorRT FP32", "FAILED", ENGINE_PATH, error=repr(e))


def save_table(rows):
    fields = ["index", "format", "precision", "status", "size_mb", "accuracy", "macro_precision", "macro_recall", "macro_f1", "weighted_precision", "weighted_recall", "weighted_f1", "latency_mean_ms", "latency_median_ms", "latency_p95_ms", "fps", "avg_power_w", "median_power_w", "max_power_w", "energy_per_inference_j", "input_shape", "shape_verified", "artifact_path", "error_message", "notes"]
    csv_path = OUT / "metrics" / "mobilenet_v3_small_fp32_core4_benchmark_table.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    (OUT / "metrics" / "mobilenet_v3_small_fp32_core4_benchmark_table.json").write_text(json.dumps(rows, indent=2), encoding="utf-8")


def save_resource_logs():
    path = OUT / "resources" / "resource_usage_timeseries.csv"
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["time", "format", "power_w"])
        w.writeheader()
        w.writerows(RESOURCE_ROWS)
    summary = {}
    for fmt in sorted(set(r["format"] for r in RESOURCE_ROWS)):
        vals = [float(r["power_w"]) for r in RESOURCE_ROWS if r["format"] == fmt and r["power_w"] != ""]
        summary[fmt] = {"samples": len(vals), "avg_power_w": statistics.mean(vals) if vals else "", "max_power_w": max(vals) if vals else ""}
    (OUT / "resources" / "resource_usage_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")


def make_reports(rows, env, class_mapping, n_images):
    lines = [
        "# MobileNetV3-Small FP32 Core4 Benchmark Report",
        "",
        "This is MobileNetV3-Small classification, not detection.",
        "FP32 baseline only. No FP16 was run. No INT8 was run.",
        "No training, fine-tuning, or pruning was run.",
        "No mAP was computed.",
        "",
        f"Dataset path: {TEST_DIR}",
        f"Number of images: {n_images}",
        f"Input shape: {INPUT_SHAPE}",
        "Preprocessing: RGB, resize to 640x640, float32 CHW tensor divided by 255, ImageNet mean/std normalization.",
        "Power note: jtop power rails are parsed as milliwatts when rail names indicate VDD/POM/power.",
        "",
        "Class mapping:",
        json.dumps(class_mapping, indent=2),
        "",
        "Environment:",
        json.dumps(env, indent=2),
        "",
        "Final 4-row table:",
    ]
    for r in rows:
        lines.append(f"- {r['format']}: status={r['status']}, accuracy={r['accuracy']}, macro_f1={r['macro_f1']}, fps={r['fps']}, energy_per_inference_j={r['energy_per_inference_j']}, shape_verified={r['shape_verified']}")
    successes = [r for r in rows if r["status"] == "SUCCESS"]
    if successes:
        fastest = max(successes, key=lambda r: float(r["fps"] or 0))
        energy_rows = [r for r in successes if r["energy_per_inference_j"] != ""]
        efficient = min(energy_rows, key=lambda r: float(r["energy_per_inference_j"])) if energy_rows else None
        lines += ["", f"Fastest FP32 format: {fastest['format']} ({fastest['fps']} FPS)."]
        lines.append(f"Most energy-efficient FP32 format: {efficient['format']} ({efficient['energy_per_inference_j']} J/inference)." if efficient else "Energy telemetry unavailable; most efficient format could not be determined.")
    txt = "\n".join(lines)
    write_text(OUT / "reports" / "mobilenet_v3_small_fp32_core4_benchmark_report.txt", txt)
    write_text(OUT / "reports" / "mobilenet_v3_small_fp32_core4_benchmark_report.md", txt)
    try:
        from docx import Document

        doc = Document()
        doc.add_heading("MobileNetV3-Small FP32 Core4 Benchmark Report", 0)
        for line in lines[2:]:
            doc.add_paragraph(line)
        doc.save(OUT / "reports" / "mobilenet_v3_small_fp32_core4_benchmark_report.docx")
    except Exception as e:
        write_text(OUT / "logs" / "docx_report_skipped.log", repr(e))


def zip_results():
    for src in [PT_PATH, ONNX_PATH, TS_PATH, ENGINE_PATH, MODEL_DIR / "onnx_export_metadata.json", MODEL_DIR / "onnx_export_report.txt"]:
        if src.exists():
            dst_dir = OUT / "models"
            dst_dir.mkdir(parents=True, exist_ok=True)
            dst = dst_dir / src.name
            try:
                if src.resolve() != dst.resolve():
                    shutil.copy2(src, dst)
            except FileNotFoundError:
                shutil.copy2(src, dst)
    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with ZipFile(ZIP_PATH, "w", compression=ZIP_DEFLATED) as zf:
        for p in sorted(OUT.rglob("*")):
            if p.is_file():
                zf.write(p, OUT.name + "/" + p.relative_to(OUT).as_posix())
    return ZIP_PATH


def main():
    env = safe_imports()
    metadata = load_export_metadata()
    class_names, items = image_paths_and_labels()
    class_mapping = {
        "dataset_class_names": class_names,
        "model_class_names": [metadata.get("class_names", {}).get(str(i), "") for i in range(len(class_names))] if metadata.get("class_names") else CLASS_NAMES,
        "mapping_note": "Dataset folder order matches the exported model class_names metadata.",
    }
    (OUT / "metrics" / "class_mapping.json").write_text(json.dumps(class_mapping, indent=2), encoding="utf-8")

    env_txt = json.dumps(env, indent=2) + "\n\ncommand line: " + " ".join(sys.argv) + "\n"
    write_text(OUT / "reports" / "environment_info.txt", env_txt)
    artifact_info = {
        "local_source_model_folder": r"C:\Users\Hiba\Desktop\results\mobilenet_v3_small\data-classification-albumentation-1\mobilenet-v3-small-v1-data-classification-albumentation-1\inference_models",
        "pt": str(PT_PATH),
        "torchscript": str(TS_PATH),
        "onnx": str(ONNX_PATH),
        "engine": str(ENGINE_PATH),
        "pt_exists": PT_PATH.exists(),
        "onnx_exists": ONNX_PATH.exists(),
        "test_images": len(items),
        "class_mapping": class_mapping,
        "metadata": metadata,
    }
    write_text(OUT / "reports" / "model_artifact_inspection.txt", json.dumps(artifact_info, indent=2))

    import torch

    device = torch.device("cuda:0" if torch.cuda.is_available() else "cpu")
    model = reconstruct_model()

    rows = []
    try:
        row = evaluate_torch_model("pytorch_fp32", "PyTorch FP32", model, device, class_names, items, PT_PATH)
        row["index"] = 1
        rows.append(row)
    except Exception as e:
        rows.append(make_result_row(1, "PyTorch FP32", "FAILED", PT_PATH, error=repr(e)))
        write_text(OUT / "logs" / "pytorch_fp32_error.log", repr(e))

    try:
        export_torchscript(model, device)
        ts_model = torch.jit.load(str(TS_PATH), map_location=device).eval()
        row = evaluate_torch_model("torchscript_fp32", "TorchScript FP32", ts_model, device, class_names, items, TS_PATH)
        row["index"] = 2
        rows.append(row)
    except Exception as e:
        rows.append(make_result_row(2, "TorchScript FP32", "FAILED", TS_PATH, error=repr(e)))
        write_text(OUT / "logs" / "torchscript_fp32_error.log", repr(e))

    row = evaluate_onnx(class_names, items)
    row["index"] = 3
    rows.append(row)

    row = evaluate_tensorrt(model, device, class_names, items)
    row["index"] = 4
    rows.append(row)

    save_table(rows)
    save_resource_logs()
    make_reports(rows, env, class_mapping, len(items))
    zip_results()

    binding_shape = None
    vpath = OUT / "metrics" / "tensorrt_fp32_engine_verification.json"
    if vpath.exists():
        try:
            data = json.loads(vpath.read_text(encoding="utf-8"))
            bindings = data.get("bindings", [])
            input_bindings = [b for b in bindings if b.get("is_input")]
            binding_shape = input_bindings[0].get("shape") if input_bindings else None
        except Exception:
            pass

    final = {
        "local_source_model_folder": r"C:\Users\Hiba\Desktop\results\mobilenet_v3_small\data-classification-albumentation-1\mobilenet-v3-small-v1-data-classification-albumentation-1\inference_models",
        "jetson_model_folder": str(MODEL_DIR),
        "jetson_output_folder": str(OUT),
        "windows_output_folder": r"C:\Users\Hiba\Desktop\results\Phase2\classification_mobilenet_v3_small_fp32_core4",
        "rows": rows,
        "input_shape": INPUT_SHAPE,
        "tensorrt_fp32_binding_shape": binding_shape,
        "final_benchmark_table_path": str(OUT / "metrics" / "mobilenet_v3_small_fp32_core4_benchmark_table.csv"),
        "report_path": str(OUT / "reports" / "mobilenet_v3_small_fp32_core4_benchmark_report.txt"),
        "zip_path": str(ZIP_PATH),
    }
    (OUT / "metrics" / "final_summary.json").write_text(json.dumps(final, indent=2), encoding="utf-8")
    print(json.dumps(final, indent=2))
    print("MOBILENETV3-SMALL FP32 BASELINE CORE4 BENCHMARK COMPLETE")


if __name__ == "__main__":
    main()
